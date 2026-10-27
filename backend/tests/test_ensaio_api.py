# Ensaio API regression tests (scripts, extract, memory, voices/TTS/STT, sessions, turns, finish)
import io
import os
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL")
if not BASE_URL:
    for ln in open("/app/frontend/.env"):
        if ln.startswith("REACT_APP_BACKEND_URL="):
            BASE_URL = ln.split("=", 1)[1].strip()
BASE_URL = BASE_URL.rstrip("/")
API = f"{BASE_URL}/api"

STRUCTURED = """TEST_ cena.

ANA: Bom dia, Pedro. Chegaste cedo hoje.
PEDRO: Não consegui dormir. Passei a noite a pensar no que me disseste.
ANA: E chegaste a alguma conclusão?
PEDRO: Cheguei. Vou aceitar o trabalho em Lisboa.
ANA: Então vamos ter de nos despedir."""


def make_pdf(text_lines):
    content = "BT /F1 12 Tf 50 750 Td 14 TL " + " ".join(f"({t}) Tj T*" for t in text_lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n".encode()
    for o in offs:
        out += f"{o:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


@pytest.fixture(scope="module")
def s():
    return requests.Session()


@pytest.fixture(scope="module")
def demo(s):
    r = s.get(f"{API}/scripts")
    assert r.status_code == 200
    d = next((x for x in r.json() if x["title"] == "A Espera"), None)
    assert d, "seeded demo missing"
    return s.get(f"{API}/scripts/{d['id']}").json()


@pytest.fixture(scope="module")
def created(s):
    r = s.post(f"{API}/scripts", json={"title": "TEST_Estruturado", "author": "t", "text": STRUCTURED}, timeout=60)
    assert r.status_code == 200, r.text
    d = r.json()
    yield d
    s.delete(f"{API}/scripts/{d['id']}")


# --- Library / seed
class TestLibrary:
    def test_root(self, s):
        assert s.get(f"{API}/").json()["message"] == "Ensaio API"

    def test_demo_seed(self, demo):
        names = [c["name"] for c in demo["characters"]]
        assert names == ["JOÃO", "MARTA", "CHEFE DA ESTAÇÃO"]
        assert "_id" not in demo
        assert sum(1 for l in demo["lines"] if l["kind"] == "dialogue") == 14

    def test_list_has_line_count(self, s):
        d = s.get(f"{API}/scripts").json()
        assert all("line_count" in x and "lines" not in x and "_id" not in x for x in d)


# --- Script create / parse
class TestScripts:
    def test_structured_parse(self, s, created):
        assert [c["name"] for c in created["characters"]] == ["ANA", "PEDRO"]
        dl = [l for l in created["lines"] if l["kind"] == "dialogue"]
        assert len(dl) == 5
        assert dl[0]["text"].startswith("Bom dia, Pedro")
        g = s.get(f"{API}/scripts/{created['id']}").json()
        assert g["title"] == "TEST_Estruturado"
        voices = {c["voice"] for c in created["characters"]}
        assert len(voices) == 2

    def test_unstructured_llm_parse(self, s):
        txt = ("O velho olhou para o mar e disse que nunca mais voltaria a pescar. A filha respondeu-lhe que "
               "isso era o que ele dizia todos os anos. Ele riu-se e disse: talvez este ano seja diferente.")
        r = s.post(f"{API}/scripts", json={"title": "TEST_Livre", "text": txt}, timeout=90)
        assert r.status_code in (200, 422), r.text
        if r.status_code == 200:
            d = r.json()
            assert len(d["characters"]) >= 1
            s.delete(f"{API}/scripts/{d['id']}")

    def test_empty_text_400(self, s):
        assert s.post(f"{API}/scripts", json={"title": "x", "text": "   "}).status_code == 400

    def test_invalid_id_404(self, s):
        assert s.get(f"{API}/scripts/nope").status_code == 404
        assert s.get(f"{API}/scripts/507f1f77bcf86cd799439011").status_code == 404

    def test_update_voice_persists(self, s, created):
        chars = [dict(c) for c in created["characters"]]
        chars[0]["voice"] = "coral"
        chars[0]["style"] = "TEST_estilo"
        r = s.put(f"{API}/scripts/{created['id']}", json={"characters": chars})
        assert r.status_code == 200
        g = s.get(f"{API}/scripts/{created['id']}").json()
        assert g["characters"][0]["voice"] == "coral" and g["characters"][0]["style"] == "TEST_estilo"


# --- Extraction
class TestExtract:
    def test_docx(self, s):
        import docx
        d = docx.Document()
        d.add_paragraph("ANA: Olá do Word.")
        d.add_paragraph("PEDRO: Resposta do Word.")
        buf = io.BytesIO()
        d.save(buf)
        r = s.post(f"{API}/extract", files={"file": ("cena.docx", buf.getvalue())})
        assert r.status_code == 200, r.text
        assert "ANA: Olá do Word." in r.json()["text"]

    def test_pdf(self, s):
        pdf = make_pdf(["ANA: Ola do PDF.", "PEDRO: Resposta do PDF."])
        r = s.post(f"{API}/extract", files={"file": ("cena.pdf", pdf)})
        assert r.status_code == 200, r.text
        assert "ANA: Ola do PDF." in r.json()["text"]

    def test_unsupported(self, s):
        r = s.post(f"{API}/extract", files={"file": ("x.exe", b"abc")})
        assert r.status_code == 400


# --- Memory
class TestMemory:
    def test_add_delete(self, s, created):
        r = s.post(f"{API}/scripts/{created['id']}/memory", json={"text": "TEST_nota"})
        assert r.status_code == 200
        nid = r.json()["id"]
        g = s.get(f"{API}/scripts/{created['id']}").json()
        assert any(m["id"] == nid and m["text"] == "TEST_nota" for m in g["memory"])
        assert s.delete(f"{API}/scripts/{created['id']}/memory/{nid}").status_code == 200
        g = s.get(f"{API}/scripts/{created['id']}").json()
        assert not any(m["id"] == nid for m in g["memory"])


# --- Voice
class TestVoice:
    def test_tts_and_stt_roundtrip(self, s):
        r = s.post(f"{API}/tts", json={"text": "Hoje é dia de pedra.", "voice": "onyx", "character": "JOÃO"}, timeout=60)
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("audio/mpeg")
        assert len(r.content) > 2000
        r2 = s.post(f"{API}/stt", files={"file": ("fala.mp3", r.content, "audio/mpeg")}, timeout=60)
        assert r2.status_code == 200, r2.text
        assert "pedra" in r2.json()["text"].lower()

    def test_tts_empty(self, s):
        assert s.post(f"{API}/tts", json={"text": "(rubrica)"}).status_code == 400

    def test_stt_tiny(self, s):
        r = s.post(f"{API}/stt", files={"file": ("a.webm", b"123")})
        assert r.status_code == 200 and r.json()["text"] == ""


# --- Sessions / turns / finish
class TestSessions:
    def test_invalid_character(self, s, demo):
        r = s.post(f"{API}/sessions", json={"script_id": demo["id"], "character": "ZE"})
        assert r.status_code == 400

    def test_full_flow(self, s, created):
        r = s.post(f"{API}/sessions", json={"script_id": created["id"], "character": "PEDRO",
                                            "instructions": "Modo rigoroso: interrompe sempre que eu não disser a fala palavra por palavra."})
        assert r.status_code == 200
        sid = r.json()["id"]
        pedro = [l for l in created["lines"] if l["character"] == "PEDRO"]
        # correct
        t = s.post(f"{API}/sessions/{sid}/turn", json={"line_idx": pedro[0]["idx"], "said": pedro[0]["text"], "mode": "escrito"}, timeout=60).json()
        assert t["turn"]["kind"] == "line" and t["advance"] is True, t
        # very wrong
        t = s.post(f"{API}/sessions/{sid}/turn", json={"line_idx": pedro[1]["idx"], "said": "Gosto muito de bananas.", "mode": "escrito"}, timeout=60).json()
        assert t["advance"] is False and t["turn"]["interrupted"] is True, t
        assert t["turn"]["feedback"]
        # note
        t = s.post(f"{API}/sessions/{sid}/turn", json={"line_idx": pedro[1]["idx"], "said": "aí mudámos nos ensaios e agora digo 'vou aceitar o emprego no Porto' em vez de 'vou aceitar o trabalho em Lisboa'", "mode": "escrito"}, timeout=60).json()
        assert t["turn"]["kind"] == "note" and t["advance"] is False and t["memory_note"], t
        # help
        t = s.post(f"{API}/sessions/{sid}/turn", json={"line_idx": pedro[1]["idx"], "said": "texto!", "mode": "escrito"}, timeout=60).json()
        assert t["turn"]["kind"] == "help" and t["advance"] is False and t["turn"]["feedback"], t
        # bad idx
        assert s.post(f"{API}/sessions/{sid}/turn", json={"line_idx": 999, "said": "x"}).status_code == 400
        # finish
        r = s.post(f"{API}/sessions/{sid}/finish", timeout=120)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["status"] == "finished"
        sm = d["summary"]
        for k in ("score", "overview", "differences", "zones_to_review", "memorization_tips", "stats"):
            assert k in sm
        assert sm["stats"]["attempts"] == 2 and sm["stats"]["interruptions"] == 1
        assert len(sm["differences"]) >= 1
        # idempotent finish
        assert s.post(f"{API}/sessions/{sid}/finish").json()["summary"]["score"] == sm["score"]
        # sessions list + script last_score
        lst = s.get(f"{API}/scripts/{created['id']}/sessions").json()
        assert any(x["id"] == sid and "turns" not in x for x in lst)
        g = s.get(f"{API}/scripts/{created['id']}").json()
        assert g["last_score"] == sm["score"]
        assert any("Porto" in m["text"] for m in g["memory"])

    def test_finish_empty_session(self, s, created):
        sid = s.post(f"{API}/sessions", json={"script_id": created["id"], "character": "ANA"}).json()["id"]
        d = s.post(f"{API}/sessions/{sid}/finish").json()
        assert d["summary"]["score"] is None
