import os
import re
import io
import json
import uuid
import hashlib
import asyncio
import logging
import tempfile
import time
from pathlib import Path
from datetime import datetime, timezone
from typing import List, Optional, Annotated, Literal

from dotenv import load_dotenv

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

from fastapi import FastAPI, APIRouter, HTTPException, UploadFile, File, Form, WebSocket  # noqa: E402
import websockets  # noqa: E402
from fastapi.responses import Response  # noqa: E402
from starlette.middleware.cors import CORSMiddleware  # noqa: E402
from motor.motor_asyncio import AsyncIOMotorClient  # noqa: E402
from pydantic import BaseModel, Field, BeforeValidator, ConfigDict  # noqa: E402
from bson import ObjectId  # noqa: E402
from emergentintegrations.llm.chat import LlmChat, UserMessage, TextDelta, StreamDone  # noqa: E402
from emergentintegrations.llm.openai import OpenAITextToSpeech, OpenAISpeechToText  # noqa: E402
import eleven  # noqa: E402
import edge_voices  # noqa: E402
import speech  # noqa: E402
import keys  # noqa: E402
import pronounce  # noqa: E402
from matching import norm_tokens, local_match, fast_verdict, final_accuracy  # noqa: E402

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("ensaio")

client = AsyncIOMotorClient(os.environ["MONGO_URL"])
db = client[os.environ["DB_NAME"]]
LLM_KEY = os.environ["EMERGENT_LLM_KEY"]
LLM_MODEL = ("anthropic", "claude-sonnet-5-5")
FAST_MODEL = ("anthropic", "claude-haiku-5-5")
TTS_CACHE = ROOT_DIR / "tts_cache"
TTS_CACHE.mkdir(exist_ok=True)

VOICES = ["onyx", "nova", "echo", "coral", "ash", "shimmer", "fable", "sage", "verse", "ballad", "alloy"]
PROMPTER_VOICE = "sage"

app = FastAPI(title="Ensaio API", openapi_url="/api/openapi.json", docs_url="/api/docs", redoc_url=None)
api = APIRouter(prefix="/api")


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def oid(v: str) -> ObjectId:
    if not ObjectId.is_valid(v):
        raise HTTPException(404, "Não encontrado")
    return ObjectId(v)


PyObjectId = Annotated[str, BeforeValidator(lambda v: str(v) if isinstance(v, ObjectId) else v)]


class BaseDocument(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    id: Optional[PyObjectId] = Field(default=None, alias="_id")

    def to_mongo(self) -> dict:
        d = self.model_dump(exclude={"id"})
        if self.id:
            d["_id"] = ObjectId(self.id)
        return d

    @classmethod
    def from_mongo(cls, doc):
        return cls.model_validate(doc) if doc else None


def short_id() -> str:
    return uuid.uuid4().hex[:10]


class Line(BaseModel):
    uid: str = Field(default_factory=short_id)
    idx: int
    character: Optional[str] = None
    text: str
    kind: str = "dialogue"


class Character(BaseModel):
    name: str
    voice: str
    style: str = ""
    el_voice_id: Optional[str] = None
    el_voice_name: str = ""
    tts: Literal["auto", "edge", "openai", "elevenlabs"] = "auto"
    edge_voice: str = ""
    edge_pitch: int = Field(0, ge=-50, le=50)
    edge_rate: int = Field(0, ge=-50, le=50)


class MemoryNote(BaseModel):
    id: str = Field(default_factory=lambda: uuid.uuid4().hex)
    text: str
    line_uid: Optional[str] = None
    new_text: Optional[str] = None
    created_at: str = Field(default_factory=now)


class Script(BaseDocument):
    title: str
    author: str = ""
    raw_text: str
    characters: List[Character] = []
    lines: List[Line] = []
    memory: List[MemoryNote] = []
    foreign: List[dict] = []
    foreign_hash: str = ""
    auto_marked: List[str] = []
    auto_marked_dismissed: bool = False
    created_at: str = Field(default_factory=now)
    last_rehearsal: Optional[str] = None
    last_score: Optional[float] = None


class Turn(BaseModel):
    id: str = Field(default_factory=short_id)
    line_idx: int
    line_idxs: List[int] = []
    review: str = "done"
    final_accuracy: Optional[float] = None
    expected: str
    said: str
    mode: str = "voz"
    kind: str = "line"
    accuracy: Optional[float] = None
    severity: str = "none"
    interrupted: bool = False
    feedback: str = ""
    differences: List[dict] = []
    created_at: str = Field(default_factory=now)


class Session(BaseDocument):
    script_id: str
    character: str
    instructions: str = ""
    strictness: str = "grave"
    start_idx: int = 0
    focus_idxs: List[int] = []
    hint_seconds: int = 5
    turns: List[Turn] = []
    status: str = "active"
    summary: Optional[dict] = None
    created_at: str = Field(default_factory=now)
    finished_at: Optional[str] = None


# ---------- LLM helpers ----------
async def ask_llm(system: str, prompt: str, model=LLM_MODEL) -> str:
    chat = LlmChat(api_key=LLM_KEY, session_id=uuid.uuid4().hex, system_message=system).with_model(*model)
    out = []
    async for ev in chat.stream_message(UserMessage(text=prompt)):
        if isinstance(ev, TextDelta):
            out.append(ev.content)
        elif isinstance(ev, StreamDone):
            break
    return "".join(out)


def parse_json(s: str):
    s = re.sub(r"```(?:json)?", "", s)
    starts = [i for i in (s.find("{"), s.find("[")) if i >= 0]
    if not starts:
        raise ValueError("sem JSON")
    start = min(starts)
    end = max(s.rfind("}"), s.rfind("]"))
    return json.loads(s[start:end + 1])


# ---------- Script parsing ----------
UP = "A-ZÁÉÍÓÚÂÊÔÃÕÇÀÜ"
RE_INLINE = re.compile(rf"^\s*([{UP}][\w{UP}à-ÿ .'\-]{{0,30}}?)\s*(\([^)]*\))?\s*[:—–]\s*(.+)$")
RE_ALONE = re.compile(rf"^\s*([{UP}][{UP}0-9 .'\-]{{0,38}}[{UP}0-9])\s*(\([^)]*\))?\s*[.:]?\s*$")


def heuristic_parse(text: str):
    raw = [r.rstrip() for r in text.replace("\r", "").split("\n")]
    alone_counts = {}
    for r in raw:
        m = RE_ALONE.match(r)
        if m:
            alone_counts[m.group(1).strip()] = alone_counts.get(m.group(1).strip(), 0) + 1
    speakers_alone = {k for k, v in alone_counts.items() if v >= 2}
    lines, cur, speaker = [], None, None

    def flush():
        nonlocal cur
        if cur and cur["text"].strip():
            lines.append(cur)
        cur = None

    for r in raw:
        s = r.strip()
        if not s:
            flush()
            speaker = None
            continue
        m_alone = RE_ALONE.match(s)
        if m_alone and m_alone.group(1).strip() in speakers_alone:
            flush()
            speaker = m_alone.group(1).strip()
            continue
        m = RE_INLINE.match(s)
        if m and len(m.group(1).split()) <= 4:
            flush()
            speaker = m.group(1).strip()
            paren = (m.group(2) + " ") if m.group(2) else ""
            cur = {"character": speaker, "text": paren + m.group(3).strip(), "kind": "dialogue"}
            continue
        if s.startswith("(") and s.endswith(")"):
            flush()
            lines.append({"character": None, "text": s, "kind": "direction"})
            continue
        if cur and cur["kind"] == "dialogue":
            cur["text"] += " " + s
        elif speaker:
            cur = {"character": speaker, "text": s, "kind": "dialogue"}
        else:
            flush()
            lines.append({"character": None, "text": s, "kind": "direction"})
    flush()
    total = sum(len(x["text"]) for x in lines) or 1
    dial = [x for x in lines if x["kind"] == "dialogue"]
    ratio = sum(len(x["text"]) for x in dial) / total
    chars = {x["character"].upper() for x in dial}
    confident = len(dial) >= 2 and ratio >= 0.6 and len(chars) >= 1
    for x in dial:
        x["character"] = x["character"].upper()
    return lines, confident


PARSE_SYSTEM = (
    "És um assistente que estrutura textos de teatro/cinema em português. Recebes um excerto e devolves "
    "APENAS JSON: uma lista de objectos {\"c\": nome da personagem em MAIÚSCULAS ou null, \"t\": texto exacto, "
    "\"k\": \"d\" para fala ou \"r\" para rubrica/indicação cénica/título}. Não alteres nem resumas o texto das falas. "
    "Se o texto for um monólogo sem nomes, atribui as falas a uma personagem com um nome adequado."
)


def chunk_text(text: str, size: int = 6000):
    parts, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        if len(cur) + len(para) > size and cur:
            parts.append(cur)
            cur = ""
        cur += para + "\n\n"
    if cur.strip():
        parts.append(cur)
    return parts


async def llm_parse(text: str):
    async def one(chunk):
        res = parse_json(await ask_llm(PARSE_SYSTEM, chunk))
        return [{"character": (x.get("c") or None) and x["c"].strip().upper(), "text": x["t"].strip(),
                 "kind": "dialogue" if x.get("k") == "d" and x.get("c") else "direction"} for x in res if x.get("t")]

    results = await asyncio.gather(*[one(c) for c in chunk_text(text)])
    return [x for r in results for x in r]


async def build_structure(text: str):
    lines, confident = heuristic_parse(text)
    if not confident:
        try:
            lines = await llm_parse(text)
        except Exception as e:
            log.exception("LLM parse falhou: %s", e)
    lines = [Line(idx=i, **x) for i, x in enumerate(lines)]
    names = []
    for ln in lines:
        if ln.kind == "dialogue" and ln.character not in names:
            names.append(ln.character)
    chars = [Character(name=n, voice=VOICES[i % len(VOICES)]) for i, n in enumerate(names)]
    return lines, chars


RE_PAGE = re.compile(r"^\s*[-–—(\[]?\s*(p[áa]g(ina)?\.?\s*)?\d{1,4}(\s*(/|de)\s*\d{1,4})?\s*[-–—)\]]?\s*$", re.I)


RE_HEADER_WORD = re.compile(r"\b(acto|ato|act|cena|quadro|parte)\b|[—–]|\s-\s", re.I)


def _printer_mark(t: str) -> bool:
    return not re.search(r"[A-Za-zÀ-ÿ]", t) and len(t) <= 20 and not t.startswith("(")


def _header_like(t: str, title: str) -> bool:
    letters = re.sub(r"[^A-Za-zÀ-ÿ]", "", t)
    upper = len(letters) >= 3 and letters == letters.upper()
    fold = lambda x: re.sub(r"\W+", " ", x.lower()).strip()  # noqa: E731
    titled = len(fold(title)) >= 4 and fold(title) in fold(t)
    author = bool(re.match(r"^(?i:de|por|texto de|by)\s+[A-ZÀ-Ý][\w'-]+(\s+[A-ZÀ-Ý][\w'-]+)+\s*$", t))
    return upper or titled or author or bool(RE_HEADER_WORD.search(t))


def _regular(pos: list) -> bool:
    gaps = [b - a for a, b in zip(pos, pos[1:])]
    return min(gaps) >= 3 and max(gaps) <= 2 * min(gaps)


def mark_non_script(lines: List[Line], title: str = "") -> List[str]:
    """Post-pass after parsing: page numbers, running headers/footers and printer marks → "foreign"."""
    page_at = [i for i, ln in enumerate(lines) if ln.kind == "direction" and (RE_PAGE.match(ln.text.strip()) or _printer_mark(ln.text.strip()))]
    near_page = lambda i: any(abs(i - p) <= 2 for p in page_at)  # noqa: E731
    near_break = lambda i: i <= 1 or i >= len(lines) - 2 or near_page(i)  # noqa: E731
    groups: dict = {}
    for i, ln in enumerate(lines):
        if ln.kind == "dialogue":
            # never digit-fold a character's line: only exact repeats of the whole "NAME — text" count
            full, k = f"{ln.character} — {ln.text.strip()}", ("d", ln.character, ln.text.strip().lower())
        elif ln.kind == "direction" and not ln.text.strip().startswith("("):
            full, k = ln.text.strip(), ("s", re.sub(r"\d+", "#", ln.text.strip().lower()))
        else:
            continue
        if len(full) <= 80:
            groups.setdefault(k, []).append((i, ln.text.strip()))
    headers = set()
    for k, occ in groups.items():
        pos = [i for i, _ in occ]
        if len(occ) < 3 or not _header_like(occ[0][1], title):
            continue
        near = sum(map(near_break, pos)) * 2 >= len(pos)
        # a character's line needs page-number evidence; edges or regular spacing alone are normal back-and-forth
        positional = sum(map(near_page, pos)) * 2 >= len(pos) if k[0] == "d" else near or _regular(pos)
        if positional:
            headers.update(pos)
    marked = []
    for i, ln in enumerate(lines):
        t = ln.text.strip()
        if i in headers:
            if ln.kind == "dialogue":
                ln.text, ln.character = f"{ln.character} — {t}", None
        elif not (ln.kind == "direction" and (RE_PAGE.match(t) or _printer_mark(t))):
            continue
        ln.kind = "foreign"
        marked.append(ln.uid)
    return marked


# ---------- Text extraction ----------
def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs)
    if name.endswith(".txt") or name.endswith(".md"):
        return data.decode("utf-8", errors="ignore")
    raise HTTPException(400, "Formato não suportado. Usa PDF, Word (.docx) ou texto (.txt).")


# ---------- Request models ----------
class ScriptIn(BaseModel):
    title: str
    author: str = ""
    text: str


class ScriptUpdate(BaseModel):
    title: Optional[str] = None
    author: Optional[str] = None
    characters: Optional[List[Character]] = None


class MemoryIn(BaseModel):
    text: str
    line_idx: Optional[int] = None
    new_text: Optional[str] = None


class LineIn(BaseModel):
    uid: Optional[str] = None
    character: Optional[str] = None
    text: str
    kind: Literal["dialogue", "direction", "foreign", "cut"] = "dialogue"


class LinesIn(BaseModel):
    lines: List[LineIn]


class RenameIn(BaseModel):
    old: str
    new: str


class TTSIn(BaseModel):
    text: str
    voice: str = PROMPTER_VOICE
    character: Optional[str] = None
    style: str = ""
    el_voice_id: Optional[str] = None
    tts: Optional[str] = None
    edge_voice: Optional[str] = None
    edge_pitch: int = 0
    edge_rate: int = 0
    script_id: Optional[str] = None


class SessionIn(BaseModel):
    script_id: str
    character: str
    instructions: str = ""
    strictness: Literal["grave", "rigor", "nunca"] = "grave"
    start_idx: int = 0
    focus_idxs: List[int] = []
    hint_seconds: int = Field(5, ge=3, le=15)


class TurnIn(BaseModel):
    line_idx: int
    line_idxs: List[int] = []
    said: str
    mode: str = "voz"


async def get_script(script_id: str) -> Script:
    s = Script.from_mongo(await db.scripts.find_one({"_id": oid(script_id)}))
    if not s:
        raise HTTPException(404, "Texto não encontrado")
    return s


async def get_session(session_id: str) -> Session:
    s = Session.from_mongo(await db.sessions.find_one({"_id": oid(session_id)}))
    if not s:
        raise HTTPException(404, "Ensaio não encontrado")
    return s


def edits_map(script: Script) -> dict:
    return {m.line_uid: m for m in script.memory if m.line_uid and m.new_text}


def effective_text(script: Script, ln: Line) -> str:
    m = edits_map(script).get(ln.uid)
    return m.new_text if m else ln.text


def script_out(script: Script) -> dict:
    d = script.model_dump()
    em = edits_map(script)
    for ln in d["lines"]:
        m = em.get(ln["uid"])
        ln["edited_text"] = m.new_text if m else None
        ln["edit_note_id"] = m.id if m else None
        ln["keyterms"] = pronounce.line_keyterms(ln["edited_text"] or ln["text"], script.foreign) if ln["kind"] == "dialogue" else []
    return d


async def save_characters(script: Script):
    await db.scripts.update_one({"_id": oid(script.id)}, {"$set": {"characters": [c.model_dump() for c in script.characters]}})
    if script.id:
        asyncio.create_task(prepare_audio(script))  # voices changed → re-pre-generate every line


LINK_SYSTEM = """Recebes as falas numeradas de um texto de teatro e uma nota de ensaio (alteração acordada). PT-PT.
Identifica a fala a que a nota se refere e escreve o texto completo da fala JÁ com a alteração aplicada (mantém tudo o resto exactamente igual).
Se a nota não for uma alteração concreta ao texto de uma fala, devolve line_idx null.
Responde APENAS JSON: {"line_idx": n|null, "new_text": "..."}"""


async def link_note(script: Script, note_text: str, hint_idx: Optional[int] = None) -> tuple:
    lines = "\n".join(f"[{ln.idx}] {ln.character}: {effective_text(script, ln)}" for ln in script.lines if ln.kind == "dialogue")
    hint = f"\nA nota foi dita durante a fala [{hint_idx}] (ou perto dela)." if hint_idx is not None else ""
    try:
        r = parse_json(await ask_llm(LINK_SYSTEM, f"Falas:\n{lines}{hint}\n\nNota: {note_text}", FAST_MODEL))
        idx = r.get("line_idx")
        if isinstance(idx, int) and 0 <= idx < len(script.lines) and r.get("new_text"):
            return script.lines[idx].uid, r["new_text"].strip()
    except Exception as e:
        log.warning("link_note falhou: %s", e)
    return None, None


async def link_note_bg(script: Script, note: MemoryNote, hint_idx: Optional[int]):
    uid, new_text = await link_note(script, note.text, hint_idx)
    if uid:
        await db.scripts.update_one({"_id": oid(script.id), "memory.id": note.id},
                                    {"$set": {"memory.$.line_uid": uid, "memory.$.new_text": new_text}})


def memory_block(script: Script) -> str:
    if not script.memory:
        return "(sem notas)"
    return "\n".join(f"- {m.text}" for m in script.memory)


# ---------- Scripts ----------
@api.get("/")
async def root():
    return {"message": "Ensaio API"}


@api.post("/extract")
async def extract(file: UploadFile = File(...)):
    data = await file.read()
    text = extract_text(file.filename or "", data).strip()
    if not text:
        raise HTTPException(400, "Não foi possível extrair texto do ficheiro.")
    return {"text": text, "filename": file.filename}


@api.get("/scripts")
async def list_scripts():
    docs = await db.scripts.find({}, {"raw_text": 0}).sort("created_at", -1).to_list(500)
    out = []
    for d in docs:
        d.setdefault("raw_text", "")
        s = Script.from_mongo(d).model_dump()
        s["line_count"] = len([x for x in s["lines"] if x["kind"] == "dialogue"])
        s.pop("lines")
        out.append(s)
    return out


@api.post("/scripts")
async def create_script(body: ScriptIn):
    if not body.text.strip():
        raise HTTPException(400, "O texto está vazio.")
    lines, chars = await build_structure(body.text)
    if not any(ln.kind == "dialogue" for ln in lines):
        raise HTTPException(422, "Não consegui identificar falas e personagens neste texto.")
    s = Script(title=body.title.strip() or "Sem título", author=body.author, raw_text=body.text, lines=lines, characters=chars)
    s.auto_marked = mark_non_script(s.lines, s.title)
    s.characters = [c for c in s.characters if any(ln.kind == "dialogue" and ln.character == c.name for ln in s.lines)]
    assign_edge(s)
    assign_eleven(s)
    res = await db.scripts.insert_one(s.to_mongo())
    s.id = str(res.inserted_id)
    spawn_foreign(s)
    return script_out(s)


@api.get("/scripts/{script_id}")
async def read_script(script_id: str):
    return script_out(await get_script(script_id))


SPOKEN_KINDS = ("dialogue", "cut")


def make_line(uid: Optional[str], idx: int, character: Optional[str], text: str, kind: str) -> Line:
    char = (character or "").strip().upper() or None
    if kind in SPOKEN_KINDS and not char:
        kind = "direction"
    return Line(uid=uid or short_id(), idx=idx, character=char if kind in SPOKEN_KINDS else None, text=text, kind=kind)


async def save_lines(script: Script, lines: List[Line]) -> dict:
    if not any(ln.kind == "dialogue" for ln in lines):
        raise HTTPException(400, "O texto tem de ter pelo menos uma fala.")
    existing = {c.name: c for c in script.characters}
    names = []
    for ln in lines:
        if ln.kind in SPOKEN_KINDS and ln.character and ln.character not in names:
            names.append(ln.character)
    used = {c.voice for c in script.characters}
    free = [v for v in VOICES if v not in used] or VOICES
    chars = [existing.get(n) or Character(name=n, voice=free[i % len(free)]) for i, n in enumerate(names)]
    script.lines, script.characters = lines, chars
    assign_edge(script)
    uids = {ln.uid for ln in lines}
    script.memory = [m for m in script.memory if not m.line_uid or m.line_uid in uids]
    await db.scripts.update_one({"_id": oid(script.id)}, {"$set": {
        "lines": [ln.model_dump() for ln in lines], "characters": [c.model_dump() for c in chars],
        "memory": [m.model_dump() for m in script.memory]}})
    spawn_foreign(script)
    return script_out(script)


@api.put("/scripts/{script_id}/lines")
async def update_lines(script_id: str, body: LinesIn):
    script = await get_script(script_id)
    lines = [x for x in body.lines if x.text.strip()]
    return await save_lines(script, [make_line(x.uid, i, x.character, x.text.strip(), x.kind) for i, x in enumerate(lines)])


class LinePatch(BaseModel):
    text: Optional[str] = None
    character: Optional[str] = None
    kind: Optional[Literal["dialogue", "direction", "foreign", "cut"]] = None


@api.patch("/scripts/{script_id}/lines/{uid}")
async def patch_line(script_id: str, uid: str, body: LinePatch):
    """Edits one line (text, character, kind) — used from the rehearsal room."""
    script = await get_script(script_id)
    ln = next((x for x in script.lines if x.uid == uid), None)
    if not ln:
        raise HTTPException(404, "Fala não encontrada")
    text = body.text.strip() if body.text is not None else ln.text
    if not text:
        raise HTTPException(400, "O texto não pode ficar vazio.")
    if body.text is not None and text != effective_text(script, ln):
        for m in script.memory:
            if m.line_uid == uid and m.new_text:
                m.line_uid, m.new_text = None, None
    kind = body.kind or ln.kind
    character = body.character if body.character is not None else ln.character
    if body.character and kind in ("direction", "foreign"):
        kind = "dialogue"
    lines = [make_line(x.uid, x.idx, character, text, kind) if x.uid == uid else x for x in script.lines]
    return await save_lines(script, lines)


@api.post("/scripts/{script_id}/auto-marked/dismiss")
async def dismiss_auto_marked(script_id: str):
    await get_script(script_id)
    await db.scripts.update_one({"_id": oid(script_id)}, {"$set": {"auto_marked_dismissed": True}})
    return script_out(await get_script(script_id))


@api.post("/scripts/{script_id}/rename-character")
async def rename_character(script_id: str, body: RenameIn):
    script = await get_script(script_id)
    new = body.new.strip().upper()
    if not new or body.old not in [c.name for c in script.characters]:
        raise HTTPException(400, "Nome inválido")
    if new == body.old:
        return script_out(script)
    for ln in script.lines:
        if ln.character == body.old:
            ln.character = new
    if new in [c.name for c in script.characters]:
        script.characters = [c for c in script.characters if c.name != body.old]
    else:
        for c in script.characters:
            if c.name == body.old:
                c.name = new
    await db.scripts.update_one({"_id": oid(script_id)}, {"$set": {
        "lines": [ln.model_dump() for ln in script.lines], "characters": [c.model_dump() for c in script.characters]}})
    return script_out(script)


@api.put("/scripts/{script_id}")
async def update_script(script_id: str, body: ScriptUpdate):
    await get_script(script_id)
    if body.characters and any(c.voice not in VOICES or (c.edge_voice and c.edge_voice not in edge_voices.ids())
                               for c in body.characters):
        raise HTTPException(400, "Voz inválida")
    upd = body.model_dump(exclude_none=True)
    if upd:
        await db.scripts.update_one({"_id": oid(script_id)}, {"$set": upd})
    out = await get_script(script_id)
    if body.characters:
        asyncio.create_task(prepare_audio(out))  # TTS settings changed → re-pre-generate every line
    return script_out(out)


@api.delete("/scripts/{script_id}")
async def delete_script(script_id: str):
    await db.scripts.delete_one({"_id": oid(script_id)})
    await db.sessions.delete_many({"script_id": script_id})
    return {"ok": True}


@api.post("/scripts/{script_id}/memory")
async def add_memory(script_id: str, body: MemoryIn):
    if not body.text.strip():
        raise HTTPException(400, "A nota está vazia.")
    script = await get_script(script_id)
    note = MemoryNote(text=body.text.strip())
    if body.line_idx is not None and body.new_text and 0 <= body.line_idx < len(script.lines):
        note.line_uid, note.new_text = script.lines[body.line_idx].uid, body.new_text.strip()
    else:
        note.line_uid, note.new_text = await link_note(script, note.text, body.line_idx)
    await db.scripts.update_one({"_id": oid(script_id)}, {"$push": {"memory": note.model_dump()}})
    return note.model_dump()


@api.delete("/scripts/{script_id}/memory/{note_id}")
async def delete_memory(script_id: str, note_id: str):
    await db.scripts.update_one({"_id": oid(script_id)}, {"$pull": {"memory": {"id": note_id}}})
    return {"ok": True}


@api.get("/scripts/{script_id}/sessions")
async def script_sessions(script_id: str):
    docs = await db.sessions.find({"script_id": script_id}).sort("created_at", -1).to_list(200)
    return [Session.from_mongo(d).model_dump(exclude={"turns"}) for d in docs]


# ---------- Voice ----------
def clean_for_tts(text: str) -> str:
    text = re.sub(r"\([^)]*\)", "", text)
    text = re.sub(r"[*_#>~|\[\]]", "", text)
    return re.sub(r"\s+", " ", text).strip()


TTS_INFLIGHT: dict = {}
TTS_SEM = asyncio.Semaphore(6)
CATALOG: dict = {"voices": [], "prompter": None}
PREMADE_IDS = {
    "JBFqnCBsd6RMkjVDRZzb", "EXAVITQu4vr4xnSDxMaL", "XrExE9yKIg1WjnnlVkGX", "nPczCjzI2devNBz1zQrb", "FGY2WhTYpPnrIDTdsKH5",
    "cjVigY5qzO86Huf0OWal", "XB0fDUnXU5powFXDhCwa", "iP95p4xoKVk53GoZ742B", "cgSgspJ2msm6clMCkdW9", "pqHfZKP75CvOlQylNhV4",
    "pFZP5JQG7iQjIQuC4Bku", "TX3LPaxmHKxFdv7VOQHJ", "Xb7hH8MSUJpSbSDYk0k2", "N2lVS1w4EtoT3dr4eOWO", "SAz9YHcvj6GT2YYXdXww",
    "onwK4e9ZLuTAKqWVaz9q",
}
TTS_STATS = {"edge": 0, "elevenlabs": 0, "openai_fallback": 0, "fallback_reasons": []}
EL_BLOCKED: set = set()  # voices the account's plan refuses (402); skipped until the key changes


def tts_instructions(character: Optional[str], style: str) -> str:
    who = f"Estás a dar voz à personagem {character}." if character else "És o apontador de um ensaio: fala de forma calma, breve e clara."
    return (
        "Fala em português de Portugal (português europeu, nunca brasileiro). "
        f"{who} {style} Diz a fala de forma natural e conversacional, como uma pessoa real numa conversa, "
        "com ritmo e respiração naturais, sem exageros nem tom de locutor."
    )


def hint_text(text: str) -> str:
    return " ".join(clean_for_tts(text).split()[:3])


async def _generate_openai(text: str, voice: str, instructions: str) -> bytes:
    engine = OpenAITextToSpeech(api_key=LLM_KEY)

    async def one(part: str) -> bytes:
        for attempt in range(2):
            try:
                return await asyncio.wait_for(engine.generate_speech(
                    text=part, model="gpt-4o-mini-tts", voice=voice, instructions=instructions), timeout=45)
            except Exception as e:
                if attempt:
                    raise
                log.warning("gpt-4o-mini-tts falhou (%s), nova tentativa", e)

    return b"".join(await asyncio.gather(*[one(p) for p in eleven.split_text(text, 1500)]))


def tts_primary() -> str:
    # User decision: Edge is the primary voice; ElevenLabs only when chosen manually for a character.
    return "edge"


def tts_cfg(c: Optional[Character]) -> dict:
    if c is None:
        pv = (CATALOG.get("prompter") or {}).get("id")
        return {"character": None, "voice": PROMPTER_VOICE, "style": "", "tts": "auto", "el_voice_id": pv, **edge_voices.PROMPTER}
    return {"character": c.name, "voice": c.voice, "style": c.style, "tts": c.tts, "el_voice_id": c.el_voice_id,
            "edge_voice": c.edge_voice or edge_voices.VOICES[0]["id"], "edge_pitch": c.edge_pitch, "edge_rate": c.edge_rate}


def _cached(path: Path, gen):
    async def run():
        if path.exists():
            return path.read_bytes()
        audio = await gen()
        path.write_bytes(audio)
        return audio
    return run


def _h(text: str, *p) -> str:
    return hashlib.sha256("|".join(map(str, (text,) + p)).encode()).hexdigest()


def _tts_edge(text: str, cfg: dict):
    ev, ep, er = cfg["edge_voice"], cfg["edge_pitch"], cfg["edge_rate"]
    return TTS_CACHE / f"ed_{_h(text, ev, ep, er)}.mp3", lambda: edge_voices.synth(text, ev, ep, er), f"{ev} {ep:+d}Hz {er:+d}%"


def _tts_eleven(text: str, cfg: dict):
    vid = cfg.get("el_voice_id")
    if not vid or vid in PREMADE_IDS or vid in EL_BLOCKED or not eleven.can("tts"):
        return None
    return TTS_CACHE / f"el_{_h(text, 'el', vid, eleven.TTS_MODEL)}.mp3", lambda: eleven.synth(text, vid), vid


def _tts_openai(text: str, cfg: dict):
    voice = cfg["voice"] if cfg["voice"] in VOICES else PROMPTER_VOICE
    instructions = tts_instructions(cfg["character"], cfg["style"])

    async def gen():
        async with TTS_SEM:
            return await _generate_openai(text, voice, instructions)
    return TTS_CACHE / f"oa_{_h(text, voice, 'gpt-4o-mini-tts', 'mp3', instructions)}.mp3", gen, voice


# Pluggable TTS: name -> builder(text, cfg) returning (cache path, generator, label) or None when unavailable.
TTS_PROVIDERS = {"elevenlabs": _tts_eleven, "edge": _tts_edge, "openai": _tts_openai}


def tts_chain(cfg: dict) -> list:
    pref = cfg.get("tts") or "auto"
    if pref == "openai":
        return ["openai"]
    first = tts_primary() if pref == "auto" else pref
    return list(dict.fromkeys([first, "edge", "openai"]))


async def _synth_job(text: str, cfg: dict) -> tuple:
    reason, tried = "", 0
    for provider in tts_chain(cfg):
        built = TTS_PROVIDERS[provider](text, cfg)
        if not built:
            continue
        path, gen, label = built
        if tried:
            log.warning("TTS FALLBACK %s (%s) texto=%r", provider, reason, text[:60])
            TTS_STATS["openai_fallback" if provider == "openai" else "fallbacks"] = TTS_STATS.get("openai_fallback" if provider == "openai" else "fallbacks", 0) + 1
            TTS_STATS["fallback_reasons"] = (TTS_STATS["fallback_reasons"] + [reason])[-20:]
        tried += 1
        try:
            audio = await _cached(path, gen)()
            TTS_STATS[provider] = TTS_STATS.get(provider, 0) + 1
            return audio, provider, label
        except Exception as e:
            reason = f"{provider}: {str(e)[:160]}"
            log.warning("TTS %s falhou: %s", provider, reason)
            if provider == "elevenlabs" and "402" in reason:
                EL_BLOCKED.add(cfg.get("el_voice_id"))
    raise HTTPException(502, "Não consegui gerar a voz.")


async def synth(text: str, cfg: dict, foreign: Optional[list] = None) -> tuple:
    text = pronounce.say_text(foreign or [], clean_for_tts(text))
    if not text:
        raise HTTPException(400, "Texto vazio")
    key = hashlib.sha256(json.dumps([text, cfg], sort_keys=True).encode()).hexdigest()
    if key not in TTS_INFLIGHT:
        async def job():
            try:
                return await _synth_job(text, cfg)
            finally:
                TTS_INFLIGHT.pop(key, None)
        TTS_INFLIGHT[key] = asyncio.ensure_future(job())
    return await asyncio.shield(TTS_INFLIGHT[key])


async def prepare_audio(script: Script, character: Optional[str] = None, start_idx: int = 0):
    """Pre-generates every AI line (start point first) so playback never waits on synthesis."""
    voices = {c.name: c for c in script.characters}
    order = script.lines[start_idx:] + script.lines[:start_idx]
    ai = [ln for ln in order if ln.kind == "dialogue" and ln.character != character and ln.character in voices]
    t0 = time.perf_counter()
    # never hold audio back for the pronunciation pass: generate now, regenerate only if respellings change
    detect = None if script.foreign_hash == foreign_hash(script) else asyncio.ensure_future(ensure_foreign(script))
    before = list(script.foreign or [])
    gen = lambda fr: asyncio.gather(*[synth(effective_text(script, ln), tts_cfg(voices[ln.character]), fr) for ln in ai], return_exceptions=True)  # noqa: E731
    res = await gen(before)
    hints = {hint_text(effective_text(script, ln)) for ln in order if character and ln.kind == "dialogue" and ln.character == character}
    await asyncio.gather(*[synth(h, tts_cfg(None)) for h in hints if h], return_exceptions=True)
    fails = [r for r in res if isinstance(r, Exception)]
    providers = [r[1] for r in res if isinstance(r, tuple)]
    log.info("pré-geração: %s falas IA (%s falhas, %s ElevenLabs, %s Edge, %s OpenAI) + %s dicas em %.1fs", len(ai), len(fails),
             providers.count("elevenlabs"), providers.count("edge"), providers.count("openai"), len(hints), time.perf_counter() - t0)
    if detect:
        try:
            after = await detect
        except Exception as e:
            log.warning("detecção de pronúncia falhou: %s", e)
            return
        if after != before:
            await gen(after)
            log.info("pré-geração refeita com %s palavras estrangeiras", len(after))


def guess_female(name: str) -> bool:
    first = name.strip().split()[0].lower() if name.strip() else ""
    return first.endswith("a") or first in {"inês", "isabel", "raquel", "beatriz", "mãe", "avó", "rainha", "senhora", "dona", "mulher"}


def assign_edge(script: Script, force: bool = False) -> bool:
    """Gives each character a distinct Edge pt-PT voice + pitch/rate combination."""
    used, changed = set(), False
    if not force:
        used = {(c.edge_voice, c.edge_pitch, c.edge_rate) for c in script.characters if c.edge_voice in edge_voices.ids()}
    for c in script.characters:
        if not force and c.edge_voice in edge_voices.ids():
            continue
        pool = edge_voices.combos("female" if guess_female(c.name) else "male")
        c.edge_voice, c.edge_pitch, c.edge_rate = next((x for x in pool if x not in used), pool[len(used) % len(pool)])
        used.add((c.edge_voice, c.edge_pitch, c.edge_rate))
        changed = True
    return changed


def assign_eleven(script: Script, force: bool = False) -> bool:
    """Distinct native pt-PT ElevenLabs voice per character (only when ElevenLabs is the primary TTS)."""
    if tts_primary() != "elevenlabs" or not CATALOG["voices"]:
        return False
    ids = {v["id"] for v in CATALOG["voices"]}
    used = set() if force else {c.el_voice_id for c in script.characters if c.el_voice_id in ids}
    changed = False
    for c in script.characters:
        if not force and c.el_voice_id in ids:
            continue
        g = "female" if guess_female(c.name) else "male"
        pool = [v for v in CATALOG["voices"] if v["gender"] == g] + [v for v in CATALOG["voices"] if v["gender"] != g]
        v = next((x for x in pool if x["id"] not in used), pool[len(used) % len(pool)])
        c.el_voice_id, c.el_voice_name = v["id"], v["name"]
        used.add(v["id"])
        changed = True
    return changed


async def eleven_ready(probe: bool = True):
    EL_BLOCKED.clear()
    if not eleven.KEY:
        return
    if probe:  # probing spends credit (TTS + STT): only on key save/check or when a character uses ElevenLabs
        await eleven.check_status(force=True)
    doc = await db.voice_catalog.find_one({"_id": "pt-PT"})
    cat = doc if doc and doc.get("voices") else eleven.seed_catalog()
    CATALOG.update(voices=cat["voices"], prompter=cat.get("prompter"), source=cat.get("source", "seed"))
    if tts_primary() == "elevenlabs":
        async for d in db.scripts.find({}):
            sc = Script.from_mongo(d)
            if assign_eleven(sc):
                await save_characters(sc)
    log.info("ElevenLabs: válida=%s, TTS principal=%s", eleven.STATUS.get("valid"), tts_primary())


async def voices_bootstrap():
    await edge_voices.refresh_voices()
    mig = await db.settings.find_one({"_id": "migrations"}) or {}
    n = 0
    async for d in db.scripts.find({}):
        s = Script.from_mongo(d)
        changed = assign_edge(s)
        if not mig.get("tts_auto"):
            for c in s.characters:
                if c.tts == "edge":
                    c.tts, changed = "auto", True
        if changed:
            await save_characters(s)
            n += 1
    await db.settings.update_one({"_id": "migrations"}, {"$set": {"tts_auto": True}}, upsert=True)
    log.info("vozes Edge: %s, textos migrados=%s", [v["id"] for v in edge_voices.VOICES], n)
    manual = await db.scripts.count_documents({"characters.tts": "elevenlabs"})
    await eleven_ready(probe=manual > 0)


def voices_info() -> dict:
    return {"provider": tts_primary(), "edge": edge_voices.VOICES, "prompter": edge_voices.PROMPTER, "openai": VOICES,
            "elevenlabs": CATALOG["voices"] if eleven.KEY else [], "eleven_configured": bool(eleven.KEY),
            "pt_native": True, "messages": []}


@api.get("/voices")
async def list_voices():
    return voices_info()


@api.get("/tts/status")
async def tts_status():
    return {"default": tts_primary(), "providers": list(TTS_PROVIDERS), "edge_voices": edge_voices.VOICES, "prompter": edge_voices.PROMPTER, "stats": TTS_STATS,
            "eleven_blocked_voices": sorted(v for v in EL_BLOCKED if v)}


@api.get("/eleven/status")
async def eleven_status(refresh: bool = False):
    st = await eleven.check_status(force=refresh)
    return {"configured": bool(eleven.KEY), "valid": st["valid"], "permissions": st["perms"], "errors": st["errors"],
            "messages": eleven.messages(), "tts_model": eleven.TTS_MODEL, "realtime_model": eleven.REALTIME_MODEL,
            "catalog": CATALOG["voices"], "prompter": CATALOG["prompter"], "stats": eleven.STATS}


@api.get("/stt/engines")
async def stt_engines():
    return {"order": ["deepgram", "gemini", "webspeech"], "deepgram": bool(speech.DEEPGRAM_KEY), "gemini": bool(speech.GEMINI_KEY),
            "deepgram_model": speech.DEEPGRAM_MODEL, "gemini_model": speech.GEMINI_LIVE_MODEL, "messages": speech.messages()}


@api.post("/stt/deepgram-token")
async def stt_deepgram_token():
    return await speech.deepgram_token()


@api.post("/stt/gemini-token")
async def stt_gemini_token():
    return await speech.gemini_token()


@api.post("/stt/token")
async def stt_token():
    if not eleven.KEY:
        raise HTTPException(503, "Falta a chave ElevenLabs no ficheiro backend/.env")
    try:
        token = await eleven.realtime_token()
    except eleven.ElevenError as e:
        log.warning("token Scribe falhou: %s", e)
        raise HTTPException(503, eleven.perm_message("realtime_token"))
    return {"token": token, "url": eleven.REALTIME_URL, "model_id": eleven.REALTIME_MODEL, "language_code": "pt", "sample_rate": 16000}


class ElevenVoiceIn(BaseModel):
    character: str
    voice: dict


@api.post("/scripts/{script_id}/eleven-voice")
async def set_eleven_voice(script_id: str, body: ElevenVoiceIn):
    if not eleven.KEY:
        raise HTTPException(400, "Falta a chave ElevenLabs no ficheiro backend/.env")
    script = await get_script(script_id)
    for c in script.characters:
        if c.name == body.character:
            c.el_voice_id, c.el_voice_name = body.voice.get("id") or None, body.voice.get("name", "")
            c.tts = "elevenlabs" if c.el_voice_id else "auto"
    if any(c.tts == "elevenlabs" for c in script.characters) and not eleven.STATUS.get("checked_at"):
        await eleven.check_status()
    await save_characters(script)
    return script_out(script)


@api.post("/scripts/{script_id}/auto-voices")
async def auto_voices(script_id: str):
    script = await get_script(script_id)
    assign_edge(script, force=True)
    assign_eleven(script, force=True)
    for c in script.characters:
        c.tts = "auto"
    await save_characters(script)
    return script_out(script)


@api.post("/tts")
async def tts(body: TTSIn):
    if body.character:
        cfg = {"character": body.character, "voice": body.voice, "style": body.style, "tts": body.tts or "auto",
               "el_voice_id": body.el_voice_id, "edge_voice": body.edge_voice if body.edge_voice in edge_voices.ids() else edge_voices.VOICES[0]["id"],
               "edge_pitch": max(-50, min(50, body.edge_pitch)), "edge_rate": max(-50, min(50, body.edge_rate))}
    else:
        cfg = tts_cfg(None)
    foreign = None
    if body.script_id and ObjectId.is_valid(body.script_id):
        d = await db.scripts.find_one({"_id": ObjectId(body.script_id)}, {"foreign": 1})
        foreign = (d or {}).get("foreign")
    audio, provider, voice = await synth(body.text, cfg, foreign)
    return Response(content=audio, media_type="audio/mpeg", headers={
        "Cache-Control": "public, max-age=86400", "X-TTS-Provider": provider, "X-TTS-Voice": voice or ""})


STT_PROMPT = "Português de Portugal. Transcreve exactamente o que é dito, palavra a palavra."


async def transcribe(data: bytes, ext: str) -> str:
    if len(data) < 1000:
        return ""
    if ext not in OpenAISpeechToText.FILE_FORMATS:
        ext = "webm"
    with tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    engine = OpenAISpeechToText(api_key=LLM_KEY)

    async def run(model: str):
        with open(tmp_path, "rb") as f:
            kw = {"prompt": STT_PROMPT} if model != "whisper-1" else {"temperature": 0}
            r = await asyncio.wait_for(engine.transcribe(file=f, model=model, language="pt", **kw), timeout=10)
        t = r.text if hasattr(r, "text") else (r.get("text") if isinstance(r, dict) else str(r))
        return (t or "").strip()

    tasks = [asyncio.ensure_future(run(m)) for m in ("gpt-4o-mini-transcribe", "whisper-1")]
    text = ""
    try:
        for fut in asyncio.as_completed(tasks):
            try:
                text = await fut
            except Exception as e:
                log.warning("STT falhou: %s", e)
                continue
            if text:
                break
    finally:
        for t in tasks:
            t.cancel()
        os.unlink(tmp_path)
    if not norm_tokens(text) or "transcreve exactamente" in " ".join(norm_tokens(text)):
        return ""
    return text


@api.post("/stt")
async def stt(file: UploadFile = File(...)):
    data = await file.read()
    return {"text": await transcribe(data, (file.filename or "a.webm").rsplit(".", 1)[-1].lower())}


# ---------- Sessions ----------
@api.post("/sessions")
async def create_session(body: SessionIn):
    script = await get_script(body.script_id)
    if body.character not in [c.name for c in script.characters]:
        raise HTTPException(400, "Personagem inválida")
    if assign_edge(script) | assign_eleven(script):
        await save_characters(script)
    s = Session(**body.model_dump())
    res = await db.sessions.insert_one(s.to_mongo())
    s.id = str(res.inserted_id)
    asyncio.create_task(prepare_audio(script, body.character, body.start_idx))
    return s.model_dump()


@api.get("/sessions/{session_id}")
async def read_session(session_id: str):
    return (await get_session(session_id)).model_dump()


TURN_SYSTEM = """És o apontador de um actor num ensaio. PT-PT. Sê rápido e breve.
Recebes: fala esperada, o que o actor disse (transcrição automática de voz ou texto escrito), instruções do actor, memória do texto (alterações acordadas que SUBSTITUEM o original) e as diferenças já detectadas automaticamente.

ATENÇÃO à transcrição automática: normaliza a fala, por isso formas coloquiais ou contraídas (prá/pra/para, tá/está, tou/estou, pró/para o, né/não é, cum/com, etc.), elisões, sotaque e pequenas palavras mal ouvidas NÃO são erros. Se a diferença pode ser explicada por pronúncia ou reconhecimento, considera que o actor disse bem.

Classifica:
- "line": tentou dizer a fala (em caso de dúvida, é "line").
- "note": informação para memorizar sobre o texto (ex.: "aí mudámos nos ensaios, agora digo X em vez de Y").
- "help": pede a deixa ("texto!", "linha", "não me lembro").
- "other": comentário claramente dirigido a ti.

Interrupção (só para "line"): segue as instruções do actor. Mesmo no modo rigoroso, interrompe apenas por palavras realmente trocadas, omitidas ou acrescentadas que mudam o texto — nunca por pronúncia, contracções, pontuação, ordem equivalente de interjeições ou possíveis erros de transcrição. Sem instruções: só erros graves (saltos grandes, sentido alterado).

Responde APENAS JSON compacto:
{"kind":"line|note|help|other","accuracy":0-100,"severity":"none|minor|major","interrupt":true|false,"feedback":"uma frase curta falada em PT-PT (se interromperes: o que falhou e o trecho correcto; se note: confirma; se help: as primeiras palavras da fala); vazio se seguir","differences":[{"expected":"","said":"","type":"omissão|troca|acrescento|repetição|salto|inversão"}],"memory_note":"para note: a nota autónoma citando a fala; senão vazio"}"""


CONFIRM_SYSTEM = """Revês diferenças entre o texto que um actor ensaiou (PT-PT) e o que um reconhecimento automático de voz (ASR) transcreveu.
O ASR é imperfeito: troca palavras raras por palavras comuns ou siglas ("arrota" → "A RTP"), corta finais ("malandreco" → "malandra"), junta e separa palavras, engole partículas e pronomes átonos, confunde homófonos, transcreve nomes estrangeiros de forma aproximada, perde palavras curtas — incluindo "não".
O actor conhece o texto. A regra por defeito é: a diferença é um ERRO DE RECONHECIMENTO.
Só é uma ALTERAÇÃO REAL do actor quando, ao mesmo tempo:
- o que foi ouvido são palavras portuguesas plausíveis que fazem sentido na frase (algo que uma pessoa diria de facto), e
- são claramente diferentes do esperado no som e no sentido (ex.: "casa" → "rua", "amanhã" → "ontem"), ou falta claramente um trecho longo (várias palavras seguidas), ou foram acrescentadas palavras com sentido próprio.
Siglas, letras soltas, palavras sem sentido no contexto, frases agramaticais, sons parecidos ou palavras curtas em falta → erro de reconhecimento.
Negações (não, nunca, nem…): só são reais se o resto da transcrição mostrar claramente que o actor inverteu o sentido; uma negação simplesmente em falta é quase sempre o ASR.
Se o modo for "escrito" não há ASR: ignora só gralhas evidentes.
Responde APENAS JSON: {"real":[índices das alterações reais], "notes":{"índice":"comentário curto em PT-PT"}}"""
CONFIRM_TASKS: dict = {}


async def confirm_differences(expected: str, said: str, diffs: List[dict], mode: str) -> tuple:
    listing = "\n".join(f"{i}. esperado: \"{d.get('expected', '')}\" → ouvido: \"{d.get('said', '')}\" ({d.get('type', '')})" for i, d in enumerate(diffs))
    prompt = f"Modo: {mode}\nFala esperada: {expected}\nTranscrição: {said}\nDiferenças:\n{listing}"
    try:
        r = parse_json(await ask_llm(CONFIRM_SYSTEM, prompt, LLM_MODEL))
        real = {int(i) for i in r.get("real", []) if str(i).lstrip("-").isdigit()}
        return real, {int(k): v for k, v in (r.get("notes") or {}).items() if str(k).isdigit()}
    except Exception as e:
        log.warning("confirmação de diferenças falhou (%s); a usar regra conservadora", e)
        return {i for i, d in enumerate(diffs) if len(d.get("expected", "").split()) >= 3}, {}


async def confirm_turn(session_id: str, turn: Turn):
    real, notes = await confirm_differences(turn.expected, turn.said, turn.differences, turn.mode)
    for i, d in enumerate(turn.differences):
        d["status"] = "real" if i in real else "discarded"
        if notes.get(i):
            d["note"] = notes[i]
    acc = final_accuracy(len(norm_tokens(turn.expected)), [d for d in turn.differences if d["status"] == "real"])
    await db.sessions.update_one({"_id": oid(session_id)}, {"$set": {
        "turns.$[t].differences": turn.differences, "turns.$[t].review": "done", "turns.$[t].final_accuracy": acc}},
        array_filters=[{"t.id": turn.id}])
    log.info("confirmação: %s/%s diferenças reais (fala %s)", len(real), len(turn.differences), turn.line_idx)


def schedule_confirm(session_id: str, turn: Turn):
    task = asyncio.create_task(confirm_turn(session_id, turn))
    tasks = CONFIRM_TASKS.setdefault(session_id, set())
    tasks.add(task)
    task.add_done_callback(tasks.discard)


def block_lines(script: Script, line_idx: int, line_idxs: List[int]) -> List[Line]:
    idxs = line_idxs or [line_idx]
    if line_idx != min(idxs) or any(i < 0 or i >= len(script.lines) for i in idxs):
        raise HTTPException(400, "Fala inválida")
    out = [script.lines[i] for i in sorted(set(idxs))]
    if any(ln.kind != "dialogue" or ln.character != out[0].character for ln in out):
        raise HTTPException(400, "Bloco inválido")
    return out


def attribute_lines(script: Script, block: List[Line], diffs: List[dict]):
    bounds, acc = [], 0
    for ln in block:
        acc += len(norm_tokens(effective_text(script, ln)))
        bounds.append((acc, ln.idx))
    for d in diffs:
        d["line_idx"] = next((idx for end, idx in bounds if d.get("pos", 0) < end), block[-1].idx)


async def evaluate_turn(session: Session, script: Script, line_idx: int, said: str, mode: str, line_idxs: List[int] = ()) -> dict:
    if line_idx < 0 or line_idx >= len(script.lines):
        raise HTTPException(400, "Fala inválida")
    block = block_lines(script, line_idx, list(line_idxs))
    line = block[0]
    expected = " ".join(effective_text(script, ln) for ln in block)
    m = local_match(expected, said, 0.9 if session.strictness == "rigor" else 0.75, pronounce.alias_pairs(script.foreign))
    attribute_lines(script, block, m["differences"])
    exp_tokens = norm_tokens(expected)
    notes = [" ".join(norm_tokens(n.text)) for n in script.memory if not n.new_text]
    in_memory = any(" ".join(exp_tokens[i:i + 4]) in n for n in notes for i in range(max(1, len(exp_tokens) - 3)))
    verdict = None if in_memory else fast_verdict(session.strictness, m, len(exp_tokens), said)
    note = None
    base = {"line_idx": line.idx, "line_idxs": [ln.idx for ln in block], "expected": expected, "said": said, "mode": mode}
    if verdict:
        _, severity = verdict
        turn = Turn(**base, kind="line", accuracy=m["accuracy"], severity=severity, differences=m["differences"])
    else:
        prev = next((ln for ln in reversed(script.lines[:line.idx]) if ln.kind == "dialogue"), None)
        prompt = (
            f"Instruções do actor: {session.instructions or '(nenhuma)'}\n"
            f"Memória do texto:\n{memory_block(script)}\n"
            f"Deixa anterior ({prev.character if prev else '-'}): {effective_text(script, prev) if prev else '-'}\n"
            f"Fala esperada ({line.character}): {expected}\n"
            f"O actor disse ({mode}): {said}\n"
            f"Diferenças automáticas, já sem as que soam igual (precisão {m['accuracy']}%): "
            f"{json.dumps([{k: d[k] for k in ('expected', 'said', 'type')} for d in m['differences']], ensure_ascii=False)}"
        )
        try:
            r = parse_json(await ask_llm(TURN_SYSTEM, prompt, FAST_MODEL))
        except Exception as e:
            log.exception("avaliação falhou: %s", e)
            r = {"kind": "line", "accuracy": m["accuracy"], "severity": "major", "interrupt": False}
        kind = r.get("kind", "line")
        turn = Turn(
            **base, kind=kind,
            accuracy=r.get("accuracy", m["accuracy"]) if kind == "line" else None, severity=r.get("severity", "none"),
            interrupted=bool(r.get("interrupt")) and kind == "line" and session.strictness != "nunca",
            feedback=r.get("feedback", "") or "",
            differences=m["differences"] if kind == "line" else [],
        )
        if kind == "note" and r.get("memory_note"):
            note = MemoryNote(text=r["memory_note"])
            await db.scripts.update_one({"_id": oid(script.id)}, {"$push": {"memory": note.model_dump()}})
            asyncio.create_task(link_note_bg(script, note, line.idx))
            if not turn.feedback:
                turn.feedback = "Anotado. Vou ter isso em conta."
    if turn.differences:
        turn.review = "pending"
        for d in turn.differences:
            d["status"] = "pending"
    elif turn.kind == "line":
        turn.final_accuracy = 100.0
    await db.sessions.update_one({"_id": oid(session.id)}, {"$push": {"turns": turn.model_dump()}})
    if turn.differences:
        schedule_confirm(session.id, turn)
    return {
        "turn": turn.model_dump(),
        "advance": turn.kind == "line" and not turn.interrupted,
        "memory_note": note.model_dump() if note else None,
    }


@api.post("/sessions/{session_id}/turn")
async def session_turn(session_id: str, body: TurnIn):
    session = await get_session(session_id)
    script = await get_script(session.script_id)
    return await evaluate_turn(session, script, body.line_idx, body.said, body.mode, body.line_idxs)


@api.post("/sessions/{session_id}/voice-turn")
async def session_voice_turn(session_id: str, line_idx: int = Form(...), file: UploadFile = File(...)):
    session = await get_session(session_id)
    script = await get_script(session.script_id)
    if line_idx < 0 or line_idx >= len(script.lines):
        raise HTTPException(400, "Fala inválida")
    data = await file.read()
    said = await transcribe(data, (file.filename or "a.webm").rsplit(".", 1)[-1].lower())
    if not said:
        return {"said": "", "turn": None, "advance": False, "memory_note": None}
    res = await evaluate_turn(session, script, line_idx, said, "voz")
    return {"said": said, **res}


def real_differences(t: Turn) -> List[dict]:
    return [d for d in t.differences if d.get("status", "real") == "real"]


async def settle_confirmations(session: Session) -> Session:
    pending = CONFIRM_TASKS.get(session.id)
    if pending:
        await asyncio.wait(list(pending), timeout=30)
    session = await get_session(session.id)
    left = [t for t in session.turns if t.review == "pending"]
    if left:
        await asyncio.gather(*[confirm_turn(session.id, t) for t in left], return_exceptions=True)
        session = await get_session(session.id)
    return session


SUMMARY_SYSTEM = """És um treinador de texto experiente para actores. Falas português de Portugal (PT-PT).
Recebes as falas que o actor ensaiou (texto esperado e, para cada tentativa, as diferenças REAIS já confirmadas — os erros de reconhecimento de voz já foram descartados), as instruções e a memória do texto.
Não inventes nem acrescentes diferenças: usa apenas as que recebes. Se não houver diferenças, diz que o texto foi dito tal e qual.
Sugere zonas a rever (só falas com diferenças reais) e dá técnicas de memorização concretas e adaptadas aos erros observados (ex.: ligar a fala à deixa, imagens mentais, palavras-âncora, ritmo).
Responde APENAS com JSON:
{"overview":"parágrafo de síntese","strengths":["..."],
 "zones_to_review":[{"line_idx":n,"reason":"..."}],
 "memorization_tips":["..."],
 "next_steps":"o que fazer no próximo ensaio"}"""


@api.post("/sessions/{session_id}/finish")
async def finish_session(session_id: str):
    session = await get_session(session_id)
    if session.status == "finished" and session.summary:
        return session.model_dump()
    session = await settle_confirmations(session)
    script = await get_script(session.script_id)
    line_turns = [t for t in session.turns if t.kind == "line"]
    attempted = sorted({i for t in line_turns for i in (t.line_idxs or [t.line_idx])})
    accs = [t.final_accuracy if t.final_accuracy is not None else (t.accuracy or 0) for t in line_turns]
    differences = [{"line_idx": d.get("line_idx", t.line_idx), "expected": d.get("expected", ""), "said": d.get("said", ""),
                    "type": d.get("type", "troca"), "note": d.get("note", "")} for t in line_turns for d in real_differences(t)]
    stats = {
        "lines_attempted": len(attempted),
        "attempts": len(line_turns),
        "interruptions": sum(1 for t in line_turns if t.interrupted),
        "avg_accuracy": round(sum(accs) / len(accs), 1) if accs else None,
        "discarded": sum(1 for t in line_turns for d in t.differences if d.get("status") == "discarded"),
    }
    if not line_turns:
        summary = {"score": None, "overview": "Não houve falas ditas neste ensaio.", "strengths": [], "differences": [],
                   "zones_to_review": [], "memorization_tips": [], "next_steps": ""}
    else:
        detail = "\n\n".join(
            f"[fala {t.line_idx}] ESPERADO: {t.expected}\nDIFERENÇAS REAIS: "
            f"{json.dumps([{k: d.get(k, '') for k in ('expected', 'said', 'type')} for d in real_differences(t)], ensure_ascii=False) if real_differences(t) else 'nenhuma'}"
            f"\n{'interrompido' if t.interrupted else 'seguiu'}"
            for t in line_turns
        )
        prompt = (f"Texto: {script.title}\nPersonagem: {session.character}\nInstruções: {session.instructions or '(nenhuma)'}\n"
                  f"Memória do texto:\n{memory_block(script)}\n\nTentativas:\n{detail}")
        try:
            summary = parse_json(await ask_llm(SUMMARY_SYSTEM, prompt))
        except Exception as e:
            log.exception("resumo falhou: %s", e)
            raise HTTPException(502, "Não consegui gerar o resumo. Tenta outra vez.")
        summary["score"] = stats["avg_accuracy"]
        summary["differences"] = differences
    summary["stats"] = stats
    fin = now()
    await db.sessions.update_one({"_id": oid(session_id)}, {"$set": {"status": "finished", "summary": summary, "finished_at": fin}})
    await db.scripts.update_one({"_id": oid(script.id)}, {"$set": {"last_rehearsal": fin, "last_score": summary.get("score")}})
    return (await get_session(session_id)).model_dump()


# ---------- Deepgram proxy (browser → backend → Deepgram) ----------
@api.websocket("/stt/deepgram")
async def stt_deepgram_ws(ws: WebSocket):
    await ws.accept()

    async def fail(msg: str, code: int):
        await ws.send_text(json.dumps({"type": "Error", "message": msg}))
        await ws.close(code)

    if not speech.DEEPGRAM_KEY:
        return await fail(speech.missing("Deepgram"), 4003)
    try:
        up = await websockets.connect(speech.deepgram_url(ws.query_params.getlist("keyterm")), open_timeout=8, max_size=None,
                                      additional_headers={"Authorization": f"Token {speech.DEEPGRAM_KEY}"})
    except Exception as e:
        code = getattr(getattr(e, "response", None), "status_code", None)
        log.warning("proxy Deepgram: ligação falhou (%s %s)", code, e)
        return await fail("Chave Deepgram inválida." if code in (401, 403) else "Sem crédito Deepgram." if code == 402
                          else f"Deepgram indisponível ({code or type(e).__name__}).", 4002)
    await ws.send_text(json.dumps({"type": "Ready"}))

    async def upstream():
        while True:
            m = await ws.receive()
            if m["type"] == "websocket.disconnect":
                return
            if m.get("bytes"):
                await up.send(m["bytes"])
            elif m.get("text"):
                await up.send(m["text"])

    async def downstream():
        async for m in up:
            await ws.send_text(m if isinstance(m, str) else m.decode())

    async def heartbeat():
        # keeps proxies/ingress from closing the socket while Deepgram is silent
        while True:
            await asyncio.sleep(15)
            await ws.send_text('{"type":"Ping"}')

    tasks = [asyncio.create_task(upstream()), asyncio.create_task(downstream()), asyncio.create_task(heartbeat())]
    await asyncio.wait(tasks, return_when=asyncio.FIRST_COMPLETED)
    for t in tasks:
        t.cancel()
    await up.close()
    try:
        await ws.close()
    except Exception:
        pass


# ---------- Definições (chaves) ----------
class KeyIn(BaseModel):
    key: str = ""


@api.get("/settings/keys")
async def settings_keys():
    return {"keys": [keys.public(p) for p in keys.ENV], "tts_primary": tts_primary()}


@api.put("/settings/keys/{provider}")
async def settings_set_key(provider: str, body: KeyIn):
    if provider not in keys.ENV:
        raise HTTPException(404, "Motor desconhecido")
    out = await keys.save(db, provider, body.key)
    if provider == "elevenlabs":
        await eleven_ready()
    return {**out, "tts_primary": tts_primary()}


@api.post("/settings/keys/{provider}/check")
async def settings_check_key(provider: str):
    if provider not in keys.ENV:
        raise HTTPException(404, "Motor desconhecido")
    k = keys.current(provider)
    if k:
        keys.STATUS[provider] = await keys.validate(provider, k)
        if provider == "elevenlabs":
            await eleven_ready()
    return {**keys.public(provider), "tts_primary": tts_primary()}


@api.post("/settings/keys/{provider}/validate")
async def settings_validate_key(provider: str, body: KeyIn):
    """Dry run: validates a candidate key without saving it or touching the key in use."""
    if provider not in keys.ENV:
        raise HTTPException(404, "Motor desconhecido")
    if not body.key.strip():
        raise HTTPException(400, "Chave vazia.")
    return {"provider": provider, "saved": False, "attempt": await keys.validate(provider, body.key.strip())}


# ---------- Pronúncia (palavras estrangeiras) ----------
FOREIGN_TASKS: dict = {}


def foreign_hash(script: Script) -> str:
    return hashlib.sha256("\n".join(effective_text(script, ln) for ln in script.lines if ln.kind == "dialogue").encode()).hexdigest()[:16]


async def ensure_foreign(script: Script, force: bool = False) -> list:
    h = foreign_hash(script)
    if not force and script.foreign_hash == h:
        return script.foreign
    if script.id in FOREIGN_TASKS and not force:
        return await FOREIGN_TASKS[script.id]

    async def run():
        text = "\n".join(f"{ln.character}: {effective_text(script, ln)}" for ln in script.lines if ln.kind == "dialogue")[:16000]
        found = pronounce.clean_items(parse_json(await asyncio.wait_for(ask_llm(pronounce.FOREIGN_SYSTEM, text), 90)).get("items"))
        body = pronounce.fold(text)
        keep = [it for it in script.foreign if pronounce.fold(it["text"]) in body]
        known = {pronounce.fold(it["text"]) for it in keep}
        items = keep + [it for it in found if pronounce.fold(it["text"]) not in known and pronounce.fold(it["text"]) in body]
        await db.scripts.update_one({"_id": oid(script.id)}, {"$set": {"foreign": items, "foreign_hash": h}})
        script.foreign, script.foreign_hash = items, h
        log.info("pronúncia: %s palavras estrangeiras em %r", len(items), script.title)
        return items

    FOREIGN_TASKS[script.id] = asyncio.ensure_future(run())
    try:
        return await FOREIGN_TASKS[script.id]
    finally:
        FOREIGN_TASKS.pop(script.id, None)


def spawn_foreign(script: Script):
    async def go():
        try:
            await ensure_foreign(script)
        except Exception as e:
            log.warning("detecção de pronúncia falhou: %s", e)
    asyncio.create_task(go())


class ForeignIn(BaseModel):
    items: List[dict]


@api.get("/scripts/{script_id}/pronunciation")
async def get_pronunciation(script_id: str):
    script = await get_script(script_id)
    return {"items": script.foreign, "detected": script.foreign_hash == foreign_hash(script)}


@api.put("/scripts/{script_id}/pronunciation")
async def put_pronunciation(script_id: str, body: ForeignIn):
    script = await get_script(script_id)
    items = pronounce.clean_items(body.items)
    await db.scripts.update_one({"_id": oid(script_id)}, {"$set": {"foreign": items, "foreign_hash": foreign_hash(script)}})
    script.foreign = items
    asyncio.create_task(prepare_audio(script))
    return {"items": items, "detected": True}


@api.post("/scripts/{script_id}/pronunciation/detect")
async def detect_pronunciation(script_id: str):
    script = await get_script(script_id)
    try:
        items = await ensure_foreign(script, force=True)
    except Exception as e:
        log.warning("detecção de pronúncia falhou: %s", e)
        raise HTTPException(502, "Não consegui detectar as palavras estrangeiras. Tenta outra vez.")
    return {"items": items, "detected": True}


DEMO_TEXT = """Cena única. Uma estação de comboios, ao fim da tarde.

(MARTA está sentada num banco. Entra JOÃO, com uma mala.)

JOÃO: Ainda estás aqui? Pensei que já tivesses partido no comboio das seis.
MARTA: Perdi-o. Ou melhor, deixei-o ir. Fiquei a vê-lo afastar-se e não me mexi.
JOÃO: Porquê?
MARTA: Porque não sabia se queria mesmo ir. Há dias em que partir parece a coisa mais fácil do mundo, e há outros em que pesa como uma pedra.
JOÃO: E hoje é dia de pedra.
MARTA: Hoje é dia de pedra.

(Pausa. JOÃO pousa a mala e senta-se ao lado dela.)

JOÃO: Sabes o que dizia a minha avó? Que as estações são os únicos sítios onde toda a gente está à espera de alguma coisa, e ninguém tem vergonha de o admitir.
MARTA: A tua avó era uma mulher sábia.
JOÃO: Era teimosa. Mas também era sábia. Uma coisa não impede a outra.
CHEFE DA ESTAÇÃO: O último comboio para Coimbra parte dentro de dez minutos, linha dois!
MARTA: Ouviste? Dez minutos.
JOÃO: Ouvi. E então? Vais ou ficas?
MARTA: Vou. Mas só se me prometeres que, quando eu voltar, ainda cá estás.
JOÃO: Prometo. Afinal, alguém tem de ficar à espera."""


@app.on_event("startup")
async def seed():
    async for d in db.scripts.find({"lines.uid": {"$exists": False}, "lines.0": {"$exists": True}}):
        lines = [{**ln, "uid": short_id()} for ln in d["lines"]]
        await db.scripts.update_one({"_id": d["_id"]}, {"$set": {"lines": lines}})
    if await db.scripts.count_documents({}) == 0:
        lines, chars = await build_structure(DEMO_TEXT)
        s = Script(title="A Espera", author="Cena de demonstração", raw_text=DEMO_TEXT, lines=lines, characters=chars)
        await db.scripts.insert_one(s.to_mongo())
    await keys.load(db)
    asyncio.create_task(voices_bootstrap())


app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["X-TTS-Provider", "X-TTS-Voice"],
)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
